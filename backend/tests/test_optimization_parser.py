import io
from types import SimpleNamespace

from docx import Document
from pptx import Presentation

from app.ingestion.parser import parse_document


def test_word_table_order_and_section():
    doc = Document()
    doc.add_heading("第一章", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "姓名", "分数"
    table.cell(1, 0).text, table.cell(1, 1).text = "Ada", "92"
    doc.add_heading("第二章", 1)
    doc.add_paragraph("其他正文")
    data = io.BytesIO()
    doc.save(data)
    parts = parse_document("a.docx", data.getvalue())
    row = next(p for p in parts if "Ada" in p.text)
    assert row.section_path == "第一章" and "姓名" in row.text
    assert parts.index(row) < next(i for i, p in enumerate(parts) if p.text == "第二章")


def test_slide_native_table():
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, 0, 0, 1000000, 1000000).table
    for cell, value in zip(
        [table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)],
        ["name", "score", "Ada", "92"],
        strict=True,
    ):
        cell.text = value
    data = io.BytesIO()
    slides.save(data)
    parts = parse_document("a.pptx", data.getvalue())
    assert any("Ada" in p.text and "score" in p.text and p.slide == 1 for p in parts)


def test_header_plus_scan_ocr_and_dedup(monkeypatch):
    import pypdf
    import pypdfium2
    import rapidocr

    pages = [SimpleNamespace(extract_text=lambda **_: "课程标题", images=[object()])]
    monkeypatch.setattr(pypdf, "PdfReader", lambda _: SimpleNamespace(pages=pages))

    class Page:
        def render(self, **kwargs):
            return SimpleNamespace(to_pil=lambda: None, close=lambda: None)

        def close(self):
            pass

    class PDF:
        def __getitem__(self, i):
            return Page()

        def close(self):
            pass

    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda _: PDF())
    monkeypatch.setattr(
        rapidocr,
        "RapidOCR",
        lambda: lambda _: SimpleNamespace(txts=["课程标题", "扫描正文 Ada 92"]),
    )
    parts = parse_document("a.pdf", b"fake")
    text = "\n".join(p.text for p in parts)
    assert text.count("课程标题") == 1 and "扫描正文 Ada 92" in text
    assert parts[0].page == 1
    monkeypatch.setattr(
        rapidocr, "RapidOCR", lambda: lambda _: SimpleNamespace(txts=[])
    )
    partial = parse_document("a.pdf", b"fake")
    assert partial.warnings and partial[0].text == "课程标题"
