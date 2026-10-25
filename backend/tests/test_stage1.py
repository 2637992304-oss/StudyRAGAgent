import io
from uuid import UUID, uuid4

import pytest
from docx import Document as WordDocument
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.auth import hash_password
from app.core.config import get_settings
from app.db.models import Chunk, Course, Document, User
from app.db.session import get_engine
from app.ingestion.parser import make_chunks, parse_document
from app.ingestion.validation import validate_upload
from app.main import app
from app.providers.local_storage import LocalStorage
from app.workers import tasks


def sample_pdf(stream: bytes = b"BT /F1 18 Tf 50 200 Td (Chapter one) Tj ET") -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + value + b"\nendobj\n")
    start = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    trailer = f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\n"
    output.extend((trailer + f"startxref\n{start}\n%%EOF").encode())
    return bytes(output)


def test_supported_file_validation() -> None:
    pdf = sample_pdf()
    word = WordDocument()
    word.add_paragraph("课程笔记")
    docx = io.BytesIO()
    word.save(docx)
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = "课程课件"
    pptx = io.BytesIO()
    slides.save(pptx)
    for filename, content in (
        ("a.pdf", pdf),
        ("a.docx", docx.getvalue()),
        ("a.pptx", pptx.getvalue()),
        ("a.txt", "课程资料".encode()),
        ("a.md", "# 课程资料".encode()),
    ):
        assert validate_upload(filename, content)
    with pytest.raises(HTTPException):
        validate_upload("fake.pdf", b"not a pdf")
    with pytest.raises(HTTPException):
        validate_upload("fake.docx", b"PK not office")
    with pytest.raises(HTTPException):
        validate_upload("empty.txt", b"")


def test_parser_keeps_slide_and_question_boundary() -> None:
    pdf_parts = parse_document("test.pdf", sample_pdf())
    assert pdf_parts[0].page == 1
    assert "Chapter one" in pdf_parts[0].text
    word = WordDocument()
    word.add_heading("第一章", level=1)
    word.add_paragraph("概念解释")
    docx = io.BytesIO()
    word.save(docx)
    word_parts = parse_document("test.docx", docx.getvalue())
    assert word_parts[-1].section_path == "第一章"
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = "第一章"
    content = io.BytesIO()
    slides.save(content)
    parts = parse_document("test.pptx", content.getvalue())
    assert make_chunks(parts)[0]["slide_no"] == 1
    parts = parse_document("test.md", "1. 问题\n答案：正确".encode())
    chunks = make_chunks(parts)
    assert [item["content_type"] for item in chunks] == ["question", "answer"]


def test_pdf_table_header_rows_and_caption_stay_in_same_context():
    content = sample_pdf(
        b"BT /F1 10 Tf 14 TL 20 250 Td "
        b"(course_id  title  dept_name  credits) Tj T* "
        b"(CS-101  Intro  Comp.Sci.  4) Tj T* "
        b"(CS-347  Databases  Comp.Sci.  3) Tj T* "
        b"(Figure A.5 The course relation.) Tj ET"
    )
    chunks = make_chunks(parse_document("university.pdf", content))
    matches = [chunk for chunk in chunks if "The course relation" in chunk["text"]]
    assert len(matches) == 1
    assert all(
        value in matches[0]["text"]
        for value in ("course_id", "title", "dept_name", "credits", "CS-101", "CS-347")
    )
    assert matches[0]["page_start"] == matches[0]["page_end"] == 1


def test_pdf_chunks_preserve_context_at_cut_without_crossing_pages():
    from app.ingestion.parser import Part

    parts = [
        Part("a" * 850 + "TABLE_HEADER " + "b" * 200, page=1),
        Part("next", page=2),
    ]
    chunks = make_chunks(parts)
    assert any("TABLE_HEADER " + "b" * 100 in chunk["text"] for chunk in chunks)
    assert all(len(chunk["text"]) <= 900 for chunk in chunks)
    assert chunks[-1]["text"] == "next" and chunks[-1]["page_start"] == 2


@pytest.fixture
def accounts(monkeypatch, tmp_path):
    settings = get_settings()
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    monkeypatch.setattr(
        tasks, "embed_texts", lambda texts: [[1.0] + [0.0] * 511 for _ in texts]
    )
    from app import api

    monkeypatch.setattr(
        api, "embed_texts", lambda texts: [[1.0] + [0.0] * 511 for _ in texts]
    )
    monkeypatch.setattr(
        api.ingest_task, "delay", lambda job_id: tasks.ingest_document(job_id)
    )
    emails = [f"{uuid4()}@example.test" for _ in range(2)]
    with Session(get_engine()) as db:
        users = [
            User(id=str(uuid4()), email=email, password_hash=hash_password("test-pass"))
            for email in emails
        ]
        db.add_all(users)
        db.commit()
        ids = [user.id for user in users]
    try:
        client = TestClient(app)
        tokens = [
            client.post(
                "/api/v1/auth/login", json={"email": email, "password": "test-pass"}
            ).json()["access_token"]
            for email in emails
        ]
        yield client, [{"Authorization": f"Bearer {token}"} for token in tokens]
    finally:
        with Session(get_engine()) as db:
            docs = db.scalars(select(Document).where(Document.owner_id.in_(ids))).all()
            for doc in docs:
                db.delete(doc)
            db.flush()
            db.execute(delete(Course).where(Course.owner_id.in_(ids)))
            db.execute(delete(User).where(User.id.in_(ids)))
            db.commit()


def test_account_isolation_ingestion_retry_and_delete(accounts, monkeypatch) -> None:
    client, (a, b) = accounts
    response = client.post("/api/v1/courses", json={"name": "数学"}, headers=a)
    assert response.status_code == 201
    course_id = response.json()["id"]
    assert client.get("/api/v1/courses", headers=b).json() == []
    assert (
        client.post(
            f"/api/v1/courses/{course_id}/search",
            headers=b,
            json={"query": "数学"},
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/courses/{course_id}/documents",
            headers=b,
            data={"document_type": "笔记"},
            files={"file": ("note.txt", b"hello", "text/plain")},
        ).status_code
        == 404
    )

    response = client.post(
        f"/api/v1/courses/{course_id}/documents",
        headers=a,
        data={"document_type": "笔记"},
        files={"file": ("note.txt", "1. 问题\n答案：正确".encode(), "text/plain")},
    )
    assert response.status_code == 202, response.text
    doc_id = response.json()["document_id"]
    job_id = response.json()["job_id"]
    with Session(get_engine()) as db:
        storage_key = UUID(db.get(Document, doc_id).storage_key)
    assert (
        client.get(f"/api/v1/documents/{doc_id}", headers=a).json()["status"] == "ready"
    )
    assert (
        client.get(f"/api/v1/jobs/{job_id}", headers=a).json()["status"] == "completed"
    )
    assert len(client.get(f"/api/v1/documents/{doc_id}/chunks", headers=a).json()) == 2
    tasks.ingest_document(job_id)
    with Session(get_engine()) as db:
        assert (
            len(db.scalars(select(Chunk).where(Chunk.document_id == doc_id)).all()) == 2
        )
    assert client.post(
        f"/api/v1/courses/{course_id}/search", headers=a, json={"query": "问题"}
    ).json()
    for path in (
        f"/api/v1/documents/{doc_id}",
        f"/api/v1/documents/{doc_id}/file",
        f"/api/v1/documents/{doc_id}/chunks",
        f"/api/v1/jobs/{job_id}",
    ):
        assert client.get(path, headers=b).status_code == 404
    assert (
        client.post(f"/api/v1/documents/{doc_id}/retry", headers=b).status_code == 404
    )
    assert client.delete(f"/api/v1/documents/{doc_id}", headers=b).status_code == 404
    assert client.delete(f"/api/v1/documents/{doc_id}", headers=a).status_code == 204
    with pytest.raises(FileNotFoundError):
        LocalStorage(get_settings().storage_dir).read(storage_key)
    tasks.ingest_document(job_id)
    assert client.get(f"/api/v1/documents/{doc_id}", headers=a).status_code == 404
    assert (
        client.post(
            f"/api/v1/courses/{course_id}/search", headers=a, json={"query": "问题"}
        ).json()
        == []
    )
    with Session(get_engine()) as db:
        assert db.scalars(select(Chunk).where(Chunk.document_id == doc_id)).all() == []


def test_delete_before_worker_runs_cannot_restore_chunks(accounts, monkeypatch) -> None:
    client, (a, _) = accounts
    course_id = client.post("/api/v1/courses", json={"name": "化学"}, headers=a).json()[
        "id"
    ]
    from app import api

    monkeypatch.setattr(api.ingest_task, "delay", lambda job_id: None)
    response = client.post(
        f"/api/v1/courses/{course_id}/documents",
        headers=a,
        data={"document_type": "笔记"},
        files={"file": ("note.md", b"# Chemistry", "text/markdown")},
    )
    doc_id = response.json()["document_id"]
    job_id = response.json()["job_id"]
    assert client.delete(f"/api/v1/documents/{doc_id}", headers=a).status_code == 204
    tasks.ingest_document(job_id)
    with Session(get_engine()) as db:
        assert db.scalars(select(Chunk).where(Chunk.document_id == doc_id)).all() == []


def test_failed_job_can_retry(accounts, monkeypatch) -> None:
    client, (a, _) = accounts
    course_id = client.post("/api/v1/courses", json={"name": "物理"}, headers=a).json()[
        "id"
    ]
    original = tasks.parse_document
    monkeypatch.setattr(
        tasks,
        "parse_document",
        lambda *_: (_ for _ in ()).throw(ValueError("解析失败")),
    )
    response = client.post(
        f"/api/v1/courses/{course_id}/documents",
        headers=a,
        data={"document_type": "笔记"},
        files={"file": ("note.txt", b"hello", "text/plain")},
    )
    doc_id = response.json()["document_id"]
    assert (
        client.get(f"/api/v1/documents/{doc_id}", headers=a).json()["status"]
        == "failed"
    )
    monkeypatch.setattr(tasks, "parse_document", original)
    response = client.post(f"/api/v1/documents/{doc_id}/retry", headers=a)
    assert response.status_code == 202
    assert (
        client.get(f"/api/v1/documents/{doc_id}", headers=a).json()["status"] == "ready"
    )
