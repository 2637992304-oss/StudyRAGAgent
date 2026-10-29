import io
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Part:
    text: str
    section_path: str | None = None
    page: int | None = None
    slide: int | None = None
    table: bool = False


class ParsedParts(list):
    version = "stage12-parse-v1"

    def __init__(self):
        super().__init__()
        self.warnings = []


class ParseError(ValueError):
    def __init__(self, warnings):
        super().__init__("文档没有可提取的文字；扫描件请检查OCR资源后重试")
        self.warnings = list(warnings)
        self.version = ParsedParts.version


def merge_ocr(native, recognized):
    lines, seen = [], set()
    for line in (native + "\n" + recognized).splitlines():
        key = " ".join(line.split()).casefold()
        if key and key not in seen:
            seen.add(key)
            lines.append(line.strip())
    return "\n".join(lines)


def parse_document(filename: str, content: bytes) -> list[Part]:
    suffix = Path(filename).suffix.lower()
    parts = ParsedParts()
    if suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        from app.providers.vision import describe_image

        parts.append(Part(text=describe_image(filename, content)))
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages, suspect = {}, []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode="layout") or ""
            if not text.strip():
                text = page.extract_text() or ""
            pages[number] = text.strip()
            images = bool(list(page.images))
            # Explainable low-text + raster heuristic, not a completeness guarantee.
            if images and len(re.sub(r"\s+", "", text)) < 160:
                suspect.append(number)
            elif images:
                parts.warnings.append(
                    f"第 {number} 页包含图像：已提取文字，图像内文字未全面核验"
                )
            elif not text.strip():
                parts.warnings.append(
                    f"第 {number} 页没有可提取文字，可能为空白页或不支持的内容"
                )
        if suspect:
            pdf = None
            try:
                import pypdfium2
                from rapidocr import RapidOCR

                ocr = RapidOCR()
                pdf = pypdfium2.PdfDocument(content)
                for number in suspect:
                    page = bitmap = image = None
                    try:
                        page = pdf[number - 1]
                        bitmap = page.render(scale=2)
                        image = bitmap.to_pil()
                        result = ocr(image)
                        recognized = "\n".join(
                            line for line in result.txts or () if line.strip()
                        )
                        if not recognized.strip():
                            raise ValueError("empty_ocr")
                        pages[number] = merge_ocr(pages[number], recognized)
                    except Exception:
                        parts.warnings.append(
                            f"第 {number} 页疑似扫描正文，OCR未完成或未识别到文字，"
                            "内容可能不完整"
                        )
                    finally:
                        for obj in (image, bitmap, page):
                            close = getattr(obj, "close", None)
                            if close:
                                close()
            except Exception:
                parts.warnings.extend(
                    f"第 {n} 页需要OCR，但OCR资源不可用，内容可能不完整"
                    for n in suspect
                )
            finally:
                if pdf is not None:
                    pdf.close()
        parts.extend(
            Part(text=text, page=number)
            for number, text in pages.items()
            if text.strip()
        )
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(io.BytesIO(content))
        section = None
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                if block.style and block.style.name.startswith("Heading"):
                    section = text
                parts.append(Part(text=text, section_path=section))
            elif isinstance(block, Table):
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in block.rows
                ]
                text = "\n".join(row for row in rows if row.strip(" |"))
                if text:
                    parts.append(Part(text=text, section_path=section, table=True))
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        for number, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                if shape.has_table:
                    rows = [
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    text = "\n".join(row for row in rows if row.strip(" |"))
                    if text:
                        parts.append(Part(text=text, slide=number, table=True))
                elif shape.has_text_frame:
                    for line in shape.text.splitlines():
                        if line.strip():
                            parts.append(Part(text=line.strip(), slide=number))
    else:
        content_text = content.decode("utf-8-sig")
        section = None
        for line in content_text.splitlines():
            if line.lstrip().startswith("#"):
                section = line.lstrip("# ").strip()
            if line.strip():
                parts.append(Part(text=line.strip(), section_path=section))
    if not parts:
        raise ParseError(parts.warnings)
    return parts


def classify(text: str) -> str:
    if re.match(r"^\s*(答案|参考答案)\s*[:：]", text):
        return "answer"
    if re.match(r"^\s*(解析|解答)\s*[:：]", text):
        return "explanation"
    if re.match(r"^\s*(第?\d+[.、)]|题目\s*[:：])", text):
        return "question"
    return "body"


def pdf_layout_units(parts: list[Part]) -> list[Part]:
    """Keep observed PDF paragraphs/DDL/table rows and pair trailing captions.

    This is text-layout preservation, not table/OCR reconstruction. No values,
    column names or relationships are invented; units never cross a page.
    """
    output = []
    for part in parts:
        if part.page is None:
            output.append(part)
            continue
        blocks = [
            "\n".join(
                " ".join(line.split()) for line in block.splitlines() if line.strip()
            )
            for block in re.split(r"\n\s*\n", part.text)
            if block.strip()
        ]
        paired = []
        for block in blocks:
            if paired and re.match(r"^Figure\s+[A-Z]?\.?\d", block):
                # Original PDF often places table captions after the data.
                paired[-1] = block + "\n" + paired[-1]
            elif paired and len(paired[-1]) < 100 and "\n" not in paired[-1]:
                # Attach a short observed heading/page heading to its next body.
                paired[-1] += "\n" + block
            else:
                paired.append(block)
        output.extend(
            Part(text=block, section_path=part.section_path, page=part.page)
            for block in paired
        )
    return output


def line_windows(
    text: str, max_chars: int, overlap: int, table: bool = False
) -> list[str]:
    """Prefer complete lines; repeat original caption/header for long tables.

    Oversized single lines use bounded character windows (recorded limitation).
    Whole-line overlap is at most the requested overlap, so always progresses.
    """
    lines = text.splitlines()
    prefix = []
    if table and lines:
        prefix = lines[:1]
    elif lines and re.match(r"^Figure\s+[A-Z]?\.?\d", lines[0]):
        prefix = lines[:2] if len(lines) > 2 else lines[:1]
    elif lines and re.match(r"^create table\b", lines[0], re.I):
        prefix = lines[:1]
    prefix_text = "\n".join(prefix)
    if len(prefix_text) > max_chars // 3:
        prefix = []
    body = lines[len(prefix) :]
    if not body:
        body, prefix = lines, []
    pieces, current = [], []
    for line in body:
        full = "\n".join(prefix + current + [line])
        if len(full) <= max_chars:
            current.append(line)
            continue
        if current:
            pieces.append("\n".join(prefix + current))
            trailing = []
            for previous in reversed(current):
                if len("\n".join([previous] + trailing)) > overlap:
                    break
                trailing.insert(0, previous)
            current = trailing
            while current and len("\n".join(prefix + current + [line])) > max_chars:
                current.pop(0)
        if len("\n".join(prefix + [line])) > max_chars:
            if current:
                current = []
            # This path does not claim to preserve an oversized row intact.
            room = max_chars - len("\n".join(prefix)) - bool(prefix)
            for start in range(0, len(line), max(1, room - min(overlap, room // 2))):
                pieces.append("\n".join(prefix + [line[start : start + room]]))
                if start + room >= len(line):
                    break
        else:
            current.append(line)
    if current:
        pieces.append("\n".join(prefix + current))
    return pieces


def make_chunks(
    parts: list[Part],
    max_chars: int = 900,
    overlap_chars: int | None = None,
    *,
    pair_headings: bool = False,
    pdf_layout: bool = False,
) -> list[dict]:
    if max_chars < 1:
        raise ValueError("max_chars 必须为正数")
    if overlap_chars is not None and not 0 <= overlap_chars < max_chars:
        raise ValueError("overlap_chars 必须大于等于零且小于 max_chars")
    chunks: list[dict] = []
    if pdf_layout:
        parts = pdf_layout_units(parts)
    if pair_headings:
        paired = []
        pending = None
        for part in parts:
            if re.fullmatch(r"\s*#{1,6}\s+[^\n]+", part.text):
                pending = part
                continue
            # Attach only to its own section/location; no source boundary crossing.
            if pending and (pending.section_path, pending.page, pending.slide) == (
                part.section_path,
                part.page,
                part.slide,
            ):
                part = Part(
                    pending.text + "\n" + part.text,
                    part.section_path,
                    part.page,
                    part.slide,
                )
            pending = None
            paired.append(part)
        parts = paired
    for part in parts:
        preserve_rows = part.table or (pdf_layout and part.page is not None)
        text = (
            "\n".join(" ".join(line.split()) for line in part.text.splitlines())
            if preserve_rows
            else " ".join(part.text.split())
        )
        kind = classify(text)
        # Keep neighboring PDF table cells together and retain context at cuts.
        overlap = (
            overlap_chars
            if overlap_chars is not None
            else (min(150, max_chars // 5) if part.page is not None else 0)
        )
        if preserve_rows:
            pieces = line_windows(text, max_chars, overlap, table=part.table)
        else:
            pieces = []
            for start in range(0, len(text), max_chars - overlap):
                pieces.append(text[start : start + max_chars])
                if start + max_chars >= len(text):
                    break
        for piece in pieces:
            if piece:
                chunks.append(
                    {
                        "text": piece,
                        "content_type": kind,
                        "section_path": part.section_path,
                        "page_start": part.page,
                        "page_end": part.page,
                        "slide_no": part.slide,
                    }
                )
    return chunks
