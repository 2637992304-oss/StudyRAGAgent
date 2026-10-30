"""Self-authored O05 comparison. SQLite exact cosine is NOT pgvector acceptance.

Run from backend: python -m scripts.evaluate_optimization_context [--live]
No paid calls without --live; all accounts, documents and files are temporary.
"""

import argparse
import hashlib
import io
import json
import math
import tempfile
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api import chat as chat_api
from app.core.auth import create_token, hash_password
from app.core.config import get_settings
from app.db.base import Base
from app.db.models import ChatSession, Chunk, Course, Document, User
from app.db.session import get_db

# Backwards-compatible name; persistent reservations are now process-safe.
from app.evaluation.request_budget import RequestBudget as Budget
from app.evaluation.runtime import observer
from app.ingestion.parser import make_chunks, parse_document
from app.main import app
from app.providers.embeddings import embed_texts
from app.providers.llm import request_budget
from app.retrieval import Candidate, PgVectorStore


def specimen(suffix, title, body):
    if suffix in {".md", ".txt"}:
        return ("# " + title + "\n" + "\n".join(body)).encode()
    stream = io.BytesIO()
    if suffix == ".docx":
        from docx import Document as Word

        doc = Word()
        doc.add_heading(title, 1)
        for text in body:
            doc.add_paragraph(text)
        doc.save(stream)
    else:
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(body)
        deck.save(stream)
    return stream.getvalue()


class ExactCosine(PgVectorStore):
    """Real local embeddings/scoped SQL; CPU cosine, no pgvector SQL claim."""

    def dense(self, vector, limit):
        norm = math.sqrt(sum(x * x for x in vector))
        items = []
        for chunk, doc in self.db.execute(self.statement()):
            other = list(chunk.embedding)
            denominator = norm * math.sqrt(sum(x * x for x in other))
            distance = (
                1 - sum(a * b for a, b in zip(vector, other, strict=True)) / denominator
                if denominator
                else 1
            )
            items.append(Candidate(chunk, doc, distance=distance))
        return sorted(items, key=lambda x: (x.distance, x.chunk.chunk_index))[:limit]


def run(live=False, output_root=None):
    import logging

    logging.getLogger("httpx").setLevel(logging.WARNING)
    settings = get_settings()
    if live and not (
        settings.llm_base_url and settings.llm_api_key and settings.llm_model
    ):
        raise RuntimeError("model_configuration_unavailable")
    root = (
        Path(output_root)
        if output_root is not None
        else Path(__file__).resolve().parents[2] / ".runtime/optimization34"
    )
    root.mkdir(parents=True, exist_ok=True)
    output = Path(tempfile.mkdtemp(prefix="context-", dir=root))
    budget = Budget(output / "budget.json")
    # Distinct from development regression topics; no tuning after this manifest.
    cases = [
        (
            ".md",
            "青石缓存",
            ["有效期为17秒。", "只有校验码相等时才允许复用。"],
            ["17秒", "校验码"],
        ),
        (
            ".txt",
            "银叶采样",
            ["每组采集23个读数。", "必须先剔除负数，再计算均值。"],
            ["23", "负数", "均值"],
        ),
        (
            ".docx",
            "紫藤调度",
            ["每批最多11个作业。", "只有上一批确认完成后才启动下一批。"],
            ["11", "完成"],
        ),
        (
            ".pptx",
            "琥珀实验",
            ["保温时间为31分钟。", "温度低于42度时本次实验无效。"],
            ["31", "42", "无效"],
        ),
    ]
    manifest = {
        "cases": cases,
        "variants": ["baseline", "short", "heading", "adjacent"],
        "top_k": settings.retrieval_top_k,
        "max_chars": settings.retrieval_max_chars,
        "model": settings.llm_model,
        "embedding": settings.embedding_model,
        "mode": "real_llm" if live else "retrieval_only",
        "store": "scoped_sqlite_exact_cosine_not_pgvector",
        "rules": "docs/optimization34.md",
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    fingerprints = {}
    # DOCX/PPTX ZIP timestamps are not a reproducible content ID: retain originals.
    sources = {
        suffix: specimen(suffix, title, body) for suffix, title, body, _ in cases
    }
    for suffix, data in sources.items():
        (output / ("source" + suffix)).write_bytes(data)
        fingerprints[suffix] = hashlib.sha256(data).hexdigest()
    results = []
    token = request_budget.set(budget)
    try:
        for variant in manifest["variants"]:
            engine = create_engine(
                "sqlite://",
                connect_args={"check_same_thread": False},
                poolclass=StaticPool,
            )
            Base.metadata.create_all(engine)
            owner, course = str(uuid4()), str(uuid4())

            def override():
                with Session(engine) as db:
                    yield db

            app.dependency_overrides[get_db] = override
            with (
                patch.object(chat_api, "get_engine", lambda: engine),
                patch("app.retrieval.PgVectorStore", ExactCosine),
                patch.object(
                    settings, "retrieval_adjacent_context", variant == "adjacent"
                ),
                patch.object(settings, "retrieval_strategy", "dense"),
                patch.object(settings, "retrieval_pair_headings", False),
            ):
                with Session(engine) as db:
                    db.add(
                        User(
                            id=owner,
                            email=owner + "@example.test",
                            password_hash=hash_password("self-authored-validation"),
                        )
                    )
                    db.add(Course(id=course, owner_id=owner, name="隔离上下文实验"))
                    db.commit()
                    for suffix, title, body, facts in cases:
                        chunks = make_chunks(
                            parse_document("source" + suffix, sources[suffix]),
                            aggregate_short=variant == "short",
                            pair_headings=variant == "heading",
                        )
                        vectors = embed_texts([c["text"] for c in chunks])
                        document = Document(
                            id=str(uuid4()),
                            owner_id=owner,
                            course_id=course,
                            filename="source" + suffix,
                            storage_key=str(uuid4()),
                            sha256=fingerprints[suffix],
                            status="ready",
                            document_type=suffix[1:],
                            mime_type="application/octet-stream",
                            size_bytes=len(sources[suffix]),
                        )
                        db.add(document)
                        db.flush()
                        for index, (chunk, vector) in enumerate(
                            zip(chunks, vectors, strict=True)
                        ):
                            db.add(
                                Chunk(
                                    id=str(uuid4()),
                                    owner_id=owner,
                                    course_id=course,
                                    document_id=document.id,
                                    chunk_index=index,
                                    content_hash=hashlib.sha256(
                                        chunk["text"].encode()
                                    ).hexdigest(),
                                    embedding_model=settings.embedding_model,
                                    parser_version="opt34-" + variant,
                                    embedding=vector,
                                    **chunk,
                                )
                            )
                        db.commit()
                        session = ChatSession(
                            owner_id=owner, course_id=course, document_ids=[document.id]
                        )
                        question = (
                            "严格根据资料，完整说明" + title + "的全部规则与限制。"
                        )
                        from time import monotonic

                        start = monotonic()
                        rows, _ = chat_api.retrieve(db, session, question)
                        text = "\n".join(c.text for c, _ in rows)
                        record = {
                            "variant": variant,
                            "case": suffix,
                            "evidence_complete": all(f in text for f in facts),
                            "required_fact_matches": [f in text for f in facts],
                            "retrieval_ms": (monotonic() - start) * 1000,
                            "chunks": len(chunks),
                            "evidence_chars": len(text),
                            "semantic_review": "not_performed",
                            "source_fingerprint": fingerprints[suffix],
                        }
                        if live:
                            db.add(session)
                            db.commit()
                            events = []
                            capture = observer.set(
                                lambda event: events.append(event)
                                if event["kind"] in {"course_sources", "model_input"}
                                else None
                            )
                            try:
                                with TestClient(app) as client:
                                    response = client.post(
                                        f"/api/v1/chat/sessions/{session.id}/messages",
                                        headers={
                                            "Authorization": "Bearer "
                                            + create_token(owner)
                                        },
                                        json={
                                            "content": question,
                                            "answer_scope": "strict",
                                            "web_enabled": False,
                                        },
                                    )
                                    record["http_status"] = response.status_code
                                    history = client.get(
                                        f"/api/v1/chat/sessions/{session.id}/messages",
                                        headers={
                                            "Authorization": "Bearer "
                                            + create_token(owner)
                                        },
                                    ).json()
                                    answer = history[-1]
                                    record.update(
                                        answer=answer["content"],
                                        status=answer["status"],
                                        measurements=answer.get("measurements"),
                                        answer_fact_matches=[
                                            f in answer["content"] for f in facts
                                        ],
                                        citations=len(answer["citations"]),
                                    )
                                    record["snapshot_ids"] = [
                                        source["chunk_id"]
                                        for event in events
                                        if event["kind"] == "course_sources"
                                        for source in event["sources"]
                                    ]
                                    record["model_input_hashes"] = [
                                        hashlib.sha256(
                                            json.dumps(
                                                event["messages"], ensure_ascii=False
                                            ).encode()
                                        ).hexdigest()
                                        for event in events
                                        if event["kind"] == "model_input"
                                    ]
                            finally:
                                observer.reset(capture)
                        results.append(record)
                        (output / "results.json").write_text(
                            json.dumps(results, ensure_ascii=False, indent=2),
                            encoding="utf-8",
                        )
            app.dependency_overrides.pop(get_db, None)
            engine.dispose()
    finally:
        request_budget.reset(token)
        app.dependency_overrides.pop(get_db, None)
    print(
        json.dumps(
            {
                "output": str(output),
                "external_requests": budget.used,
                "rows": len(results),
                "adopted": False,
                "reason": "small_sample_requires_semantic_review",
            }
        )
    )
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--live",
        action="store_true",
        help="Up to 60 actual external requests using existing credentials",
    )
    run(parser.parse_args().live)
