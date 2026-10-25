"""Run 20 real retrieval and LLM questions against temporary public samples.

Requires PostgreSQL, Redis, a Celery worker, and a configured LLM. Run from
backend/ with `uv run python scripts/smoke_stage2.py`.
"""

import io
import json
import time
from pathlib import Path
from uuid import uuid4

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.auth import create_token, hash_password
from app.core.config import get_settings
from app.db.models import ChatSession, Course, User
from app.db.session import get_engine
from app.main import app

SAMPLES = [
    (
        "教材",
        "textbook.pdf",
        [
            (
                "What is photosynthesis?",
                "Photosynthesis converts light energy into chemical energy.",
            ),
            (
                "What does chlorophyll absorb?",
                "Chlorophyll absorbs light for photosynthesis.",
            ),
            (
                "What gas is released in photosynthesis?",
                "Photosynthesis releases oxygen.",
            ),
        ],
    ),
    (
        "课件",
        "slides.pptx",
        [
            ("什么是细胞膜？", "细胞膜控制物质进出细胞。"),
            ("线粒体的作用是什么？", "线粒体通过呼吸作用提供能量。"),
            ("细胞核保存什么？", "细胞核保存遗传信息。"),
        ],
    ),
    (
        "笔记",
        "notes.txt",
        [
            ("导数表示什么？", "导数表示函数在一点的瞬时变化率。"),
            ("积分可用于计算什么？", "定积分可以计算曲线下的有向面积。"),
            ("极限描述什么？", "极限描述变量趋近某个值时函数的趋势。"),
        ],
    ),
    (
        "试卷",
        "exam.md",
        [
            ("三角形内角和是多少？", "三角形内角和是 180 度。"),
            (
                "直角三角形的勾股关系是什么？",
                "勾股关系是两条直角边平方和等于斜边平方。",
            ),
            ("圆周率通常记作什么？", "圆周率通常记作 π。"),
        ],
    ),
    (
        "作业习题",
        "homework.docx",
        [
            ("水的化学式是什么？", "水的化学式是 H2O。"),
            ("二氧化碳的化学式是什么？", "二氧化碳的化学式是 CO2。"),
            ("食盐的主要化学成分是什么？", "食盐的主要成分是氯化钠。"),
        ],
    ),
]


def pdf_bytes(lines: list[str]) -> bytes:
    stream = (
        "BT /F1 11 Tf 40 740 Td "
        + " ".join(
            f"({line.replace('(', '').replace(')', '')}) Tj 0 -20 Td" for line in lines
        )
        + " ET"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    start = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        (
            f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\n"
            f"startxref\n{start}\n%%EOF"
        ).encode()
    )
    return bytes(data)


def make_file(filename: str, facts: list[tuple[str, str]]) -> bytes:
    lines = [fact for _, fact in facts]
    if filename.endswith(".pdf"):
        return pdf_bytes(lines)
    if filename.endswith(".pptx"):
        slides = Presentation()
        for line in lines:
            slide = slides.slides.add_slide(slides.slide_layouts[1])
            slide.shapes.title.text = "课程知识点"
            slide.placeholders[1].text = line
        output = io.BytesIO()
        slides.save(output)
        return output.getvalue()
    if filename.endswith(".docx"):
        word = WordDocument()
        for line in lines:
            word.add_paragraph(line)
        output = io.BytesIO()
        word.save(output)
        return output.getvalue()
    return "\n".join(lines).encode()


def parse_events(response) -> list[tuple[str, dict]]:
    return [
        (block.split("\n")[0][7:], json.loads(block.split("data: ", 1)[1]))
        for block in response.text.strip().split("\n\n")
        if "data: " in block
    ]


def wait_ready(client: TestClient, doc_id: str, headers: dict) -> None:
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        response = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
        response.raise_for_status()
        status = response.json()["status"]
        if status == "ready":
            return
        if status == "failed":
            raise RuntimeError(f"资料处理失败：{doc_id}")
        time.sleep(1)
    raise TimeoutError(f"资料处理超时：{doc_id}")


def main() -> None:
    users = [str(uuid4()), str(uuid4())]
    with Session(get_engine()) as db, db.begin():
        db.add_all(
            User(
                id=uid,
                email=f"smoke2-{uid}@example.test",
                password_hash=hash_password("test"),
            )
            for uid in users
        )
    client = TestClient(app)
    a, b = [{"Authorization": f"Bearer {create_token(uid)}"} for uid in users]
    rows = []
    doc_ids = []
    try:
        course_id = client.post(
            "/api/v1/courses", json={"name": "阶段二公开样本"}, headers=a
        ).json()["id"]
        for doc_type, filename, facts in SAMPLES:
            response = client.post(
                f"/api/v1/courses/{course_id}/documents",
                headers=a,
                data={"document_type": doc_type},
                files={"file": (filename, make_file(filename, facts))},
            )
            response.raise_for_status()
            doc_id = response.json()["document_id"]
            doc_ids.append(doc_id)
            wait_ready(client, doc_id, a)
        for index, (doc_type, filename, facts) in enumerate(SAMPLES):
            session = client.post(
                "/api/v1/chat/sessions",
                json={"course_id": course_id, "document_ids": [doc_ids[index]]},
                headers=a,
            )
            session.raise_for_status()
            session_id = session.json()["id"]
            for question, expected_fact in facts:
                response = client.post(
                    f"/api/v1/chat/sessions/{session_id}/messages",
                    json={"content": question},
                    headers=a,
                )
                result = parse_events(response)
                citations = [data for name, data in result if name == "citation"]
                terminal = [data for name, data in result if name in {"done", "error"}]
                passed = (
                    response.status_code == 200
                    and any(
                        citation["document_id"] == doc_ids[index]
                        for citation in citations
                    )
                    and any(name == "done" for name, _ in result)
                )
                rows.append(
                    {
                        "question": question,
                        "category": doc_type,
                        "scope": [doc_ids[index]],
                        "expected_document": filename,
                        "expected_fact": expected_fact,
                        "answer": "".join(
                            data["text"]
                            for name, data in result
                            if name == "answer_delta"
                        ),
                        "actual_documents": [item["filename"] for item in citations],
                        "actual_chunk_ids": [item["chunk_id"] for item in citations],
                        "actual_locations": [
                            item["page_start"] or item["slide_no"] for item in citations
                        ],
                        "source_check_passed": passed,
                        "terminal_event": result[-1][0] if result else None,
                        "terminal_status": terminal[-1].get("status")
                        if terminal
                        else None,
                        "error_message": terminal[-1].get("message")
                        if terminal
                        else None,
                    }
                )
                time.sleep(3)
        # Five boundary cases complete the 20-case set.
        empty_course = client.post(
            "/api/v1/courses", json={"name": "无资料课程"}, headers=a
        ).json()["id"]
        empty_session = client.post(
            "/api/v1/chat/sessions", json={"course_id": empty_course}, headers=a
        ).json()["id"]
        for question in ("这门课的教材讲了什么？", "未上传资料时能回答吗？"):
            result = parse_events(
                client.post(
                    f"/api/v1/chat/sessions/{empty_session}/messages",
                    json={"content": question},
                    headers=a,
                )
            )
            passed = not any(name == "citation" for name, _ in result) and any(
                data.get("status") == "no_evidence"
                for name, data in result
                if name == "done"
            )
            rows.append(
                {
                    "question": question,
                    "category": "无资料",
                    "scope": [],
                    "expected_document": None,
                    "answer": "".join(
                        data["text"] for name, data in result if name == "answer_delta"
                    ),
                    "actual_documents": [],
                    "actual_chunk_ids": [],
                    "actual_locations": [],
                    "source_check_passed": passed,
                }
            )
        own_session = client.post(
            "/api/v1/chat/sessions",
            json={"course_id": course_id, "document_ids": [doc_ids[0]]},
            headers=a,
        ).json()["id"]
        for question in ("未选中的笔记如何定义导数？", "其他课程的内容是什么？"):
            result = parse_events(
                client.post(
                    f"/api/v1/chat/sessions/{own_session}/messages",
                    json={"content": question},
                    headers=a,
                )
            )
            citations = [data for name, data in result if name == "citation"]
            rows.append(
                {
                    "question": question,
                    "category": "资料范围",
                    "scope": [doc_ids[0]],
                    "expected_document": None,
                    "answer": "".join(
                        data["text"] for name, data in result if name == "answer_delta"
                    ),
                    "actual_documents": [item["filename"] for item in citations],
                    "actual_chunk_ids": [item["chunk_id"] for item in citations],
                    "actual_locations": [item["page_start"] for item in citations],
                    "source_check_passed": all(
                        item["document_id"] == doc_ids[0] for item in citations
                    ),
                }
            )
        forbidden = client.post(
            f"/api/v1/chat/sessions/{own_session}/messages",
            json={"content": "跨账号读取"},
            headers=b,
        )
        rows.append(
            {
                "question": "账号 B 读取账号 A 的会话",
                "category": "跨账号",
                "scope": [own_session],
                "expected_document": None,
                "answer": "",
                "actual_documents": [],
                "actual_chunk_ids": [],
                "actual_locations": [],
                "source_check_passed": forbidden.status_code == 404,
            }
        )
        assert len(rows) == 20
        output = Path("../docs/stage2_acceptance_results.json")
        output.write_text(
            json.dumps(
                {
                    "llm_model": get_settings().llm_model,
                    "embedding_model": get_settings().embedding_model,
                    "cases": rows,
                    "passed": sum(item["source_check_passed"] for item in rows),
                    "answer_factuality_reviewed": False,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        passed = sum(item["source_check_passed"] for item in rows)
        print(f"Stage 2 source checks: {passed}/20")
        print(f"Results: {output}")
    finally:
        with Session(get_engine()) as db:
            for chat in db.scalars(
                select(ChatSession).where(ChatSession.owner_id.in_(users))
            ).all():
                db.delete(chat)
            db.commit()
        for doc_id in doc_ids:
            client.delete(f"/api/v1/documents/{doc_id}", headers=a)
        with Session(get_engine()) as db:
            for course in db.scalars(
                select(Course).where(Course.owner_id.in_(users))
            ).all():
                db.delete(course)
            for user in db.scalars(select(User).where(User.id.in_(users))).all():
                db.delete(user)
            db.commit()


if __name__ == "__main__":
    main()
